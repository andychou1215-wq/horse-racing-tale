"""從 `數值平衡試算.xlsx` 重新產生 tests/ 使用的參考資料 data/reference_horses.json。

當 xlsx 的「屬性設定」或「比賽模擬」分頁內容有更新時，重新執行本腳本即可
更新測試基準值。

用法：
    python3 tools/regen_reference.py /path/to/數值平衡試算.xlsx
"""
from __future__ import annotations

import json
import pathlib
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "data" / "reference_horses.json"


def main(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    ws_attr = wb["屬性設定"]
    headers = [c.value for c in ws_attr[1]]
    horses = []
    for row in ws_attr.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            break
        horses.append(dict(zip(headers, row)))

    ws_race = wb["比賽模擬"]
    race_headers = [c.value for c in ws_race[4]]
    race_rows = []
    for row in ws_race.iter_rows(min_row=5, max_row=4 + len(horses), values_only=True):
        race_rows.append(dict(zip(race_headers, row)))

    meta = {
        "distance_mod": ws_race["B1"].value,
        "terrain_mod": ws_race["B2"].value,
    }

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUTPUT, "w", encoding="utf-8") as f:
        json.dump(
            {"meta": meta, "horses": horses, "race_results": race_rows},
            f,
            ensure_ascii=False,
            indent=2,
        )
    print(f"寫入 {OUTPUT}（{len(horses)} 匹測試馬）")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("用法: python3 tools/regen_reference.py /path/to/數值平衡試算.xlsx")
        sys.exit(1)
    main(sys.argv[1])
